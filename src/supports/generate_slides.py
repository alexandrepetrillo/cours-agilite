#!/usr/bin/env python3
"""
Génération des slides PPTX pour la journée complète
Méthodologies Agile — Journée 2 (09h00 – 17h00)
Slides allégés : titres + bullet points essentiels uniquement
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Palette de couleurs ---
BLUE_DARK = RGBColor(0x00, 0x3F, 0x5C)
BLUE_ACCENT = RGBColor(0x00, 0x77, 0xB6)
BLUE_LIGHT = RGBColor(0x90, 0xE0, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_DARK = RGBColor(0x33, 0x33, 0x33)
GRAY_LIGHT = RGBColor(0x66, 0x66, 0x66)
ORANGE = RGBColor(0xFF, 0x6B, 0x35)
GREEN = RGBColor(0x2D, 0x9C, 0x4A)
RED = RGBColor(0xE0, 0x3E, 0x3E)


def add_bg_shape(slide, color=BLUE_DARK):
    """Ajoute un rectangle de fond coloré en haut de la slide."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), Inches(1.2)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # Envoyer en arrière-plan
    sp = shape._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)


def add_section_slide(prs, title, subtitle="", emoji=""):
    """Slide de section (fond coloré complet)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank

    # Fond complet
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BLUE_DARK
    bg_shape.line.fill.background()
    sp = bg_shape._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)

    # Emoji grand
    if emoji:
        txBox = slide.shapes.add_textbox(Inches(0), Inches(1.5), Inches(13.33), Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = emoji
        run.font.size = Pt(72)

    # Titre
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(3.0), Inches(12.33), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = WHITE

    # Sous-titre
    if subtitle:
        txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11.33), Inches(1))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(22)
        run2.font.color.rgb = BLUE_LIGHT
        run2.font.italic = True

    return slide


def add_title_slide(prs, title, subtitle=""):
    """Slide avec titre et bullet points."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank

    add_bg_shape(slide)

    # Titre dans la bande
    txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(12), Inches(0.9))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = WHITE

    if subtitle:
        txBox2 = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(12), Inches(0.5))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(16)
        run2.font.color.rgb = GRAY_LIGHT
        run2.font.italic = True

    return slide


def add_bullets(slide, items, left=0.6, top=2.0, width=12, font_size=20, spacing=Pt(8)):
    """Ajoute des bullet points à une slide."""
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(5)
    )
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        p.space_after = spacing

        # Support pour sous-items (tuple: (text, level))
        if isinstance(item, tuple):
            text, level = item
            p.level = level
        else:
            text = item
            p.level = 0

        # Support pour texte en gras partiel (séparé par **)
        if "**" in text:
            parts = text.split("**")
            for j, part in enumerate(parts):
                if part:
                    run = p.add_run()
                    run.text = part
                    run.font.size = Pt(font_size - (2 * p.level))
                    run.font.color.rgb = GRAY_DARK
                    if j % 2 == 1:  # Les parties impaires sont en gras
                        run.font.bold = True
                        run.font.color.rgb = BLUE_ACCENT
        else:
            run = p.add_run()
            run.text = text
            run.font.size = Pt(font_size - (2 * p.level))
            run.font.color.rgb = GRAY_DARK

    return txBox


def add_two_columns(slide, left_items, right_items, left_title="", right_title="", top=2.0):
    """Ajoute deux colonnes de texte."""
    # Titre colonne gauche
    if left_title:
        txBox = slide.shapes.add_textbox(Inches(0.6), Inches(top - 0.4), Inches(5.5), Inches(0.4))
        tf = txBox.text_frame
        run = tf.paragraphs[0].add_run()
        run.text = left_title
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = BLUE_ACCENT

    if right_title:
        txBox = slide.shapes.add_textbox(Inches(7), Inches(top - 0.4), Inches(5.5), Inches(0.4))
        tf = txBox.text_frame
        run = tf.paragraphs[0].add_run()
        run.text = right_title
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = BLUE_ACCENT

    add_bullets(slide, left_items, left=0.6, top=top, width=5.5, font_size=18)
    add_bullets(slide, right_items, left=7, top=top, width=5.5, font_size=18)


def add_highlight_box(slide, text, left=2, top=5.5, width=9, color=ORANGE):
    """Ajoute un encadré de mise en avant."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.8)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = WHITE


# ============================================================
#  CONSTRUCTION DU DECK
# ============================================================

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# ============================================================
#  SLIDE 1 — COUVERTURE
# ============================================================
add_section_slide(
    prs,
    "Méthodologies Agile",
    "Journée 2 — US, Kanban, XP & Clôture  |  09h00 – 17h00",
    emoji="🚀"
)

# ============================================================
#  SLIDE 2 — PROGRAMME DE LA JOURNÉE
# ============================================================
s = add_title_slide(prs, "📋  Programme de la journée")
add_bullets(s, [
    "**09h00** — Stand-Up Comedy (Daily revisité)",
    "**09h15** — User Stories & Definition of Done",
    "**09h45** — Atelier US + priorisation + chiffrage",
    "**10h45** — Ball Point Game (vélocité)",
    "**11h30** — Restaurant Kanban (flux, WIP, Lead Time)",
    "**13h15** — Cours XP + Live Coding TDD",
    "**13h55** — Pair Programming & TDD (kata)",
    "**14h45** — Refactoring Race",
    "**15h45** — Sprint Review & Rétrospective",
    "**16h20** — Jeopardy Agile + Clôture",
], font_size=20)

# ============================================================
#  SECTION — STAND-UP
# ============================================================
add_section_slide(prs, "Stand-Up Comedy", "09h00 – 09h15  |  Daily Stand-up revisité", "🎤")

s = add_title_slide(prs, "🎤  Stand-Up Comedy — Règles")
add_bullets(s, [
    "Debout en cercle, bâton de parole",
    "**30 secondes** max par personne",
    "",
    "🟢  « La dernière fois j'ai appris que… »",
    "🔵  « Aujourd'hui j'attends… »",
    "🔴  « Mon obstacle c'est… »",
])
add_highlight_box(s, "C'est un vrai Daily Scrum — 15 min debout, 3 questions !")

# ============================================================
#  SECTION — USER STORIES & DoD
# ============================================================
add_section_slide(prs, "User Stories & Definition of Done", "09h15 – 09h45  |  Cours interactif", "📝")

s = add_title_slide(prs, "📝  Format d'une User Story")
add_bullets(s, [
    "« **En tant que** [persona] »",
    "« **je veux** [fonctionnalité] »",
    "« **afin de** [bénéfice/valeur] »",
    "",
    "Exemple :",
    ("**En tant qu'**étudiant, **je veux** consulter le menu de la cantine, **afin de** choisir mon repas avant d'y aller", 1),
])
add_highlight_box(s, "Une US décrit un BESOIN utilisateur, pas une solution technique")

s = add_title_slide(prs, "✅  Critères INVEST")
add_bullets(s, [
    "**I**ndependent — pas de dépendance entre US",
    "**N**egotiable — discutable, pas gravée dans le marbre",
    "**V**aluable — apporte de la valeur à l'utilisateur",
    "**E**stimable — l'équipe peut estimer l'effort",
    "**S**mall — réalisable en 1 sprint",
    "**T**estable — on peut vérifier que c'est fait",
])

s = add_title_slide(prs, "🏁  Definition of Done (DoD)")
add_bullets(s, [
    "Sans DoD, « **terminé** » ne veut rien dire",
    "",
    "Exemple de DoD :",
    ("Code écrit + tests unitaires passent", 1),
    ("Code reviewé par un pair", 1),
    ("Documentation mise à jour", 1),
    ("Déployé en environnement de recette", 1),
    ("Critères d'acceptation validés par le PO", 1),
    "",
    "**DoD** = checklist commune à TOUTES les US",
    "**Critères d'acceptation** = spécifiques à UNE US",
])

# ============================================================
#  SECTION — ATELIER US
# ============================================================
add_section_slide(prs, "Atelier US + Priorisation + Chiffrage", "09h45 – 10h30  |  Pizza Campus 🍕", "🔧")

s = add_title_slide(prs, "🍕  Brief — Application Pizza Campus")
add_bullets(s, [
    "Contexte : app de commande de pizzas pour le campus",
    "",
    "**Étape 1** — Rédiger **8-10 User Stories** (15 min)",
    ("Identifier les personas d'abord (étudiant, pizzaïolo, admin…)", 1),
    ("Format : En tant que… je veux… afin de…", 1),
    ("Sur post-its, 1 US par post-it", 1),
    "",
    "**Étape 2** — Priorisation **MoSCoW** (5 min)",
    ("🔴 Must  |  🟡 Should  |  🟢 Could  |  ⚪ Won't", 1),
    "",
    "**Étape 3** — **Planning Poker** sur les Must (15 min)",
    ("Fibonacci : 1, 2, 3, 5, 8, 13, 21", 1),
])

# ============================================================
#  SECTION — BALL POINT GAME
# ============================================================
add_section_slide(prs, "Ball Point Game", "10h45 – 11h30  |  Vélocité & Estimation", "🏐")

s = add_title_slide(prs, "🏐  Ball Point Game — Règles")
add_bullets(s, [
    "Toute la classe en cercle",
    "Chaque balle doit passer par **toutes les mains**",
    "Pas de passage au **voisin direct**",
    "Balle qui tombe = **-1 point**  |  Balle complétée = **+1 point**",
    "",
    "**5 sprints de 2 min** avec estimation + rétro entre chaque",
    "",
    "Avant chaque sprint : **« Combien de balles pensez-vous réussir ? »**",
])
add_highlight_box(s, "Estimation → Exécution → Mesure → Rétro → Amélioration = le cycle Agile")

s = add_title_slide(prs, "📈  Ball Point Game — Débrief")
add_bullets(s, [
    "**Vélocité** = nombre de points réalisés par sprint",
    "**Estimation** = prédiction avant le sprint",
    "",
    "On trace la **courbe de vélocité** sprint après sprint",
    "→ la vélocité se **stabilise** après 2-3 sprints",
    "→ on peut **prédire** le sprint suivant",
    "",
    "**Burn-up chart** = cumul des points sprint après sprint",
    "→ visualise la progression vers l'objectif",
])

# ============================================================
#  SECTION — KANBAN
# ============================================================
add_section_slide(prs, "Kanban — Le Restaurant", "11h30 – 12h15  |  Simulation physique", "🍽️")

s = add_title_slide(prs, "🍽️  L'analogie Restaurant → IT")
add_two_columns(s,
    [
        "Les commandes clients",
        "La cuisine",
        "Commande → Prépa → Cuisson → Dressage → Service",
        "Le nombre de feux",
        "Temps commande → service",
        "Trop de commandes en même temps",
    ],
    [
        "Les tickets / tâches",
        "Le workflow de développement",
        "Backlog → Analyse → Dev → Test → Done",
        "La **limite WIP**",
        "Le **Lead Time**",
        "**Pas de limite WIP** → chaos !",
    ],
    left_title="🍽️ Restaurant",
    right_title="📊 Kanban / Projet IT",
    top=2.2,
)

s = add_title_slide(prs, "📊  Les 5 principes Kanban")
add_bullets(s, [
    "1. **Visualiser** le flux de travail",
    "2. **Limiter** le travail en cours (WIP)",
    "3. **Gérer** le flux",
    "4. **Rendre explicites** les règles",
    "5. **Améliorer** de manière collaborative",
])
add_highlight_box(s, "« Stop starting, start finishing. »")

s = add_title_slide(prs, "🍽️  Simulation — Le Menu")
add_bullets(s, [
    "🍔  **Burger simple** — 1 étoile — 3 ingrédients à dessiner",
    "🍕  **Pizza** — 2 étoiles — 5 ingrédients à dessiner",
    "🥘  **Plat gastronomique** — 3 étoiles — 8 ingrédients à dessiner",
    "",
    "**Round 1** : PAS de limite WIP → chaos garanti",
    "**Round 2** : Limite WIP = 2 par poste → flux tiré",
    "",
    "Chaque cuisinier **annonce à voix haute** le ticket",
    "Le Maître d'hôtel fait un **SNAPSHOT** chaque minute",
])

s = add_title_slide(prs, "📈  Kanban — Métriques clés")
add_bullets(s, [
    "**Lead Time** — temps total du début à la fin",
    "**Cycle Time** — temps en traitement actif",
    "**Throughput** — nombre d'items terminés par unité de temps",
    "**Cumulative Flow Diagram** — aires empilées montrant les encours",
    "",
    "Round 1 vs Round 2 :",
    ("Throughput ≈ pareil (le goulot ne change pas !)", 1),
    ("Lead Time bien meilleur en R2", 1),
    ("Moins de stress, moins d'erreurs", 1),
    ("Board Kanban fiable en R2", 1),
])
add_highlight_box(s, "Kanban ne supprime pas l'attente — il la rend visible et honnête")

# ============================================================
#  SECTION — XP
# ============================================================
add_section_slide(prs, "Extreme Programming (XP)", "13h15 – 13h55  |  Cours + Live Coding", "💻")

s = add_title_slide(prs, "💻  XP — Les 5 valeurs")
add_bullets(s, [
    "💬  **Communication** — parler, pas documenter en silo",
    "🎯  **Simplicité** — faire le plus simple qui fonctionne",
    "🔄  **Feedback** — boucles courtes, apprendre vite",
    "💪  **Courage** — refactorer, jeter, dire non",
    "🤝  **Respect** — confiance mutuelle dans l'équipe",
])

s = add_title_slide(prs, "🧰  XP — Les pratiques clés")
add_two_columns(s,
    [
        "👥  **Pair Programming**",
        "→ Pilote + Navigateur",
        "→ 2 cerveaux > 1",
        "",
        "🧪  **TDD**",
        "→ Red → Green → Refactor",
        "→ Le test AVANT le code",
    ],
    [
        "🔄  **Intégration Continue**",
        "→ Commit souvent, testez toujours",
        "",
        "🧹  **Refactoring**",
        "→ Améliorer sans changer le comportement",
        "",
        "📋  **Planning Game**",
        "→ Estimation collective",
    ],
    top=2.0,
)

s = add_title_slide(prs, "🧪  TDD — Le cycle Red-Green-Refactor")
add_bullets(s, [
    "🔴  **RED** — Écrire un test qui ÉCHOUE",
    ("Le test décrit le comportement attendu", 1),
    "",
    "🟢  **GREEN** — Écrire le code MINIMUM pour faire passer le test",
    ("Pas d'optimisation, pas de gold-plating", 1),
    "",
    "🔵  **REFACTOR** — Améliorer le code SANS casser les tests",
    ("Nommage, extraction de fonctions, simplification", 1),
    "",
    "→ Recommencer avec le prochain test",
])
add_highlight_box(s, "Live Coding : kata FizzBuzz en TDD")

# ============================================================
#  SECTION — PAIR PROGRAMMING
# ============================================================
add_section_slide(prs, "Pair Programming & TDD", "13h55 – 14h45  |  Kata Roman Numerals", "👥")

s = add_title_slide(prs, "🏓  Ping-Pong TDD — Règles")
add_bullets(s, [
    "En **binômes** (tirage au sort)",
    "",
    "🏓  Joueur A écrit un **test qui échoue** (RED)",
    "🏓  Joueur B écrit le **code minimum** pour le faire passer (GREEN)",
    "🏓  Joueur B écrit le **prochain test** qui échoue (RED)",
    "🏓  Joueur A écrit le **code** pour le faire passer (GREEN)",
    "",
    "Rotation toutes les **2-3 minutes**",
    "",
    "Progression : 1→I, 2→II, 3→III, 4→IV, 5→V, 9→IX, 10→X…",
])

# ============================================================
#  SECTION — REFACTORING RACE
# ============================================================
add_section_slide(prs, "Refactoring Race", "14h45 – 15h30  |  Compétition chronométrée", "🏁")

s = add_title_slide(prs, "🏁  Refactoring Race — Règles")
add_bullets(s, [
    "Même **code spaghetti** pour toutes les équipes",
    "**20 minutes** chrono — musique de fond 🎵",
    "Objectif : refactorer **sans casser les tests**",
    "",
    "**Scoring :**",
    ("✅  Tests passent toujours : **+10 pts**", 1),
    ("📝  Nommage clair : **+5 pts**", 1),
    ("🧩  Fonctions < 20 lignes : **+5 pts**", 1),
    ("🔁  Pas de duplication : **+5 pts**", 1),
    ("🧪  Nouveaux tests : **+3 pts/test**", 1),
    ("💡  SOLID : **+5 pts bonus**", 1),
])

# ============================================================
#  SECTION — SPRINT REVIEW & RETRO
# ============================================================
add_section_slide(prs, "Sprint Review & Rétrospective", "15h45 – 16h20", "🎬")

s = add_title_slide(prs, "🎬  Sprint Review — Théorie")
add_bullets(s, [
    "**Quoi ?** Démo de l'incrément aux stakeholders",
    "**Qui ?** Équipe + PO + parties prenantes",
    "**Quand ?** Fin de chaque sprint",
    "**Résultat ?** Feedback → adaptation du Product Backlog",
    "",
    "Rappel Kapla : « Les démos entre les sprints = des Sprint Reviews ! »",
])
add_highlight_box(s, "Review = inspecter le PRODUIT  |  Rétro = inspecter le PROCESSUS")

s = add_title_slide(prs, "⭐  Rétrospective — Étoile de Mer")
add_bullets(s, [
    "5 branches :",
    "",
    "▶️  **Continuer** — ce qui a bien fonctionné",
    "➕  **Plus de** — ce qu'on voudrait davantage",
    "🆕  **Commencer** — ce qu'on devrait initier",
    "➖  **Moins de** — ce qu'on devrait réduire",
    "⏹️  **Arrêter** — ce qui ne fonctionne pas",
    "",
    "1 post-it par branche → lecture collective → **actions concrètes**",
])

# ============================================================
#  SECTION — JEOPARDY
# ============================================================
add_section_slide(prs, "Jeopardy Agile + Clôture", "16h20 – 17h00", "🏆")

s = add_title_slide(prs, "🏆  Jeopardy Agile — Catégories")
add_bullets(s, [
    "5 catégories × 3 niveaux (100 / 200 / 300 points)",
    "",
    "📜  **Manifeste Agile**",
    "🏉  **Scrum**",
    "📊  **Kanban**",
    "💻  **XP**",
    "🌍  **Agile IRL** (cas pratiques)",
    "",
    "30 secondes pour répondre — mauvaise réponse = vol possible !",
])

s = add_title_slide(prs, "🏆  Trophées & Clôture")
add_bullets(s, [
    "🥇  **Équipe la plus Agile** — meilleur score global",
    "💡  **Question la plus pertinente** de la formation",
    "🏗️  **Plus belle construction Kapla**",
    "",
    "**ROTI** — Return On Time Invested (vote 1→5)",
    "**3 mots** pour résumer la formation → nuage de mots",
    "",
    "📚  Ressources :",
    ("Scrum.org  |  Kanbanize.com  |  Agile Alliance", 1),
    ("Vidéo : « Spotify Engineering Culture »", 1),
    ("Vidéo : « Agile Product Ownership in a Nutshell »", 1),
])
add_highlight_box(s, "Merci et bonne continuation ! 🚀", color=GREEN)

# ============================================================
#  SLIDE FINALE
# ============================================================
add_section_slide(
    prs,
    "Merci !",
    "« Les individus et leurs interactions\nplus que les processus et les outils »\n— Manifeste Agile, 2001",
    emoji="🙏"
)

# --- SAUVEGARDE ---
output_path = "/home/racing/workspace/perso/ort-agilite-plan/src/supports/slides-journee-complete.pptx"
prs.save(output_path)
print(f"✅ Slides générés : {output_path}")
print(f"   {len(prs.slides)} slides au total")

